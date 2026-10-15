from flask import Flask, render_template, request, jsonify, send_file

from transformers import AutoTokenizer, AutoModelForSeq2SeqLM

from PyPDF2 import PdfReader

from fpdf import FPDF

from docx import Document

from docx.shared import Pt

from werkzeug.utils import secure_filename

from structure import (
    extract_sections_from_pdf,
    analyze_text_structure,
    sections_to_text
)

import os
import re
import torch


# ============================================================
# FLASK APPLICATION
# ============================================================

app = Flask(__name__)


# ============================================================
# CONFIGURATION
# ============================================================

MODEL_PATH = "./academic_text_translator_model"

MAX_DOCUMENT_WORDS = 10000

OUTPUT_FOLDER = "generated_files"

os.makedirs(
    OUTPUT_FOLDER,
    exist_ok=True
)


# ============================================================
# TRANSLATION SPEED SETTINGS
# ============================================================

# Translate several chunks together.
# This is the main speed improvement.
if torch.cuda.is_available():
    BATCH_SIZE = 8
else:
    BATCH_SIZE = 4

# Faster decoding than the old 5-beam search.
NUM_BEAMS = 2

# Input and output limits.
MAX_INPUT_TOKENS = 400
MAX_NEW_TOKENS = 400


# ============================================================
# LOAD TRANSLATION MODEL
# ============================================================

print("Loading translation model...")

tokenizer = AutoTokenizer.from_pretrained(
    MODEL_PATH
)

model = AutoModelForSeq2SeqLM.from_pretrained(
    MODEL_PATH
)

# Automatically use GPU if available.
if torch.cuda.is_available():

    DEVICE = torch.device("cuda")

else:

    DEVICE = torch.device("cpu")

model.to(DEVICE)

model.eval()

print(
    "Translation model loaded successfully!"
)

print(
    "Using device:",
    DEVICE
)

print(
    "Translation batch size:",
    BATCH_SIZE
)


# ============================================================
# FIND HINDI FONT
# ============================================================

def find_hindi_font():

    possible_fonts = [

        # Windows Nirmala UI font collection
        r"C:\Windows\Fonts\Nirmala.ttc",

        # Other possible Hindi fonts
        r"C:\Windows\Fonts\Mangal.ttf",

        r"C:\Windows\Fonts\Nirmala.ttf",

        # Project fonts folder
        "./fonts/Nirmala.ttc",

        "./fonts/Nirmala.ttf",

        "./fonts/Mangal.ttf",
    ]

    for font in possible_fonts:

        if os.path.exists(font):

            print(
                "Hindi font found:",
                font
            )

            return font

    print(
        "No Hindi font found."
    )

    return None


HINDI_FONT = find_hindi_font()


def find_hindi_bold_font():

    possible_fonts = [

        r"C:\Windows\Fonts\NirmalaB.ttf",

        r"C:\Windows\Fonts\mangalb.ttf",

        "./fonts/NirmalaB.ttf",

        "./fonts/mangalb.ttf",
    ]

    for font in possible_fonts:

        if os.path.exists(font):

            print("Hindi bold font found:", font)

            return font

    print(
        "No Hindi bold font found - "
        "headings will use colour and larger size."
    )

    return None


HINDI_BOLD_FONT = find_hindi_bold_font()


# ============================================================
# DOCUMENT STRUCTURE
# ============================================================
# Extraction and structure detection live in structure.py:
#   extract_sections_from_pdf(bytes) -> (sections, text)
#   analyze_text_structure(text)     -> sections


# ============================================================
# SPLIT TEXT INTO CHUNKS (sentence aware)
# ============================================================

def split_into_chunks(
    text,
    max_words=180
):

    sentences = re.split(
        r"(?<=[.!?])\s+",
        text.strip()
    )

    chunks = []

    current = []

    current_words = 0

    for sentence in sentences:

        count = len(sentence.split())

        # A single very long sentence: cut it by words.
        if count > max_words:

            if current:
                chunks.append(" ".join(current))
                current = []
                current_words = 0

            words = sentence.split()

            for i in range(0, len(words), max_words):
                chunks.append(
                    " ".join(words[i:i + max_words])
                )

            continue

        if current_words + count > max_words:

            chunks.append(" ".join(current))

            current = []

            current_words = 0

        current.append(sentence)

        current_words += count

    if current:
        chunks.append(" ".join(current))

    return [c for c in chunks if c.strip()]


# ============================================================
# TRANSLATE A BATCH
# ============================================================

def translate_batch(texts):

    if not texts:

        return []


    # Tokenize several chunks together.
    inputs = tokenizer(

        texts,

        return_tensors="pt",

        padding=True,

        truncation=True,

        max_length=MAX_INPUT_TOKENS

    )


    # Move tensors to CPU/GPU.
    inputs = {

        key: value.to(DEVICE)

        for key, value in inputs.items()

    }


    # No gradient calculation is needed
    # because this is inference only.
    with torch.inference_mode():

        outputs = model.generate(

            **inputs,

            # Faster than the old num_beams=5.
            num_beams=NUM_BEAMS,

            do_sample=False,

            max_new_tokens=MAX_NEW_TOKENS,

            no_repeat_ngram_size=3,

            early_stopping=True

        )


    translations = tokenizer.batch_decode(

        outputs,

        skip_special_tokens=True

    )


    return [
        translation.strip()
        for translation in translations
    ]


# ============================================================
# TRANSLATE STRUCTURED DOCUMENT
# ============================================================
# Every section keeps its type (title / subtitle / heading /
# paragraph / bullet / number), its heading level and its number.
# Only the "text" is translated.

def prepare_for_translation(section):

    text = section["text"]

    # ALL-CAPS headings translate badly: make them normal case first.
    if section["type"] in ("title", "heading", "subtitle") \
            and text.isupper():

        text = text.title()

    return text


def translate_document(sections):

    jobs = []

    for section_index, section in enumerate(sections):

        for chunk_index, chunk in enumerate(
            split_into_chunks(
                prepare_for_translation(section),
                max_words=180
            )
        ):

            jobs.append({
                "section_index": section_index,
                "chunk_index": chunk_index,
                "text": chunk
            })

    total_jobs = len(jobs)

    print(
        f"Translating {total_jobs} chunks "
        f"in batches of {BATCH_SIZE}..."
    )

    results = {}

    for start in range(0, total_jobs, BATCH_SIZE):

        batch = jobs[start:start + BATCH_SIZE]

        translations = translate_batch(
            [job["text"] for job in batch]
        )

        for job, translation in zip(batch, translations):

            results.setdefault(
                job["section_index"],
                []
            ).append(
                (job["chunk_index"], translation)
            )

        print(
            f"Translated "
            f"{min(start + BATCH_SIZE, total_jobs)}"
            f"/{total_jobs} chunks"
        )

    translated_sections = []

    for section_index, section in enumerate(sections):

        parts = sorted(results.get(section_index, []))

        if not parts:
            continue

        translated = dict(section)   # keeps type, level, number

        translated["text"] = " ".join(
            text for _, text in parts
        ).strip()

        translated_sections.append(translated)

    return translated_sections


# ============================================================
# COUNT WORDS
# ============================================================

def count_words(text):

    return len(

        re.findall(

            r"\b\w+\b",

            text

        )

    )


# ============================================================
# STYLE TABLE (same look for PDF output)
# ============================================================

HEADING_COLOR = (26, 61, 124)

# type/level -> (font size, bold, colour, space before, space after)
def pdf_style(section):

    kind = section["type"]

    if kind == "title":
        return 20, True, (0, 0, 0), 2, 4

    if kind == "subtitle":
        return 11, False, (110, 110, 110), 0, 8

    if kind == "heading":

        level = section.get("level", 2)

        if level == 1:
            return 15, True, HEADING_COLOR, 6, 3

        if level == 2:
            return 13, True, HEADING_COLOR, 4, 2

        return 12, True, HEADING_COLOR, 3, 2

    return 11, False, (0, 0, 0), 0, 2


# ============================================================
# GENERATE HINDI PDF
# ============================================================

def generate_pdf(
    sections,
    original_filename
):

    if not HINDI_FONT:

        raise RuntimeError(
            "Hindi font not found. "
            "Please make sure Nirmala UI or "
            "Mangal is installed on Windows."
        )

    output_name = (
        os.path.splitext(
            secure_filename(original_filename)
        )[0]
        + "_Hindi_Translation.pdf"
    )

    output_path = os.path.join(
        OUTPUT_FOLDER,
        output_name
    )

    pdf = FPDF()

    pdf.set_auto_page_break(auto=True, margin=18)

    pdf.set_margins(18, 18, 18)

    pdf.add_page()

    # --------------------------------------------------------
    # FONTS (regular + bold)
    # --------------------------------------------------------

    if HINDI_FONT.lower().endswith(".ttc"):

        pdf.add_font(
            family="Hindi",
            fname=HINDI_FONT,
            collection_font_number=0
        )

    else:

        pdf.add_font(
            family="Hindi",
            fname=HINDI_FONT
        )

    has_bold = False

    if HINDI_BOLD_FONT:

        try:

            pdf.add_font(
                family="Hindi",
                style="B",
                fname=HINDI_BOLD_FONT
            )

            has_bold = True

        except Exception as error:

            print("Bold font could not be loaded:", error)

    pdf.set_text_shaping(
        True,
        script="deva",
        language="hin"
    )

    # --------------------------------------------------------
    # If the document has its own title, mirror it exactly.
    # Otherwise print the default banner.
    # --------------------------------------------------------

    has_title = any(
        section["type"] == "title"
        for section in sections
    )

    if not has_title:

        pdf.set_font("Hindi", size=18)

        pdf.cell(
            0,
            12,
            text="शैक्षणिक पाठ का हिंदी अनुवाद",
            new_x="LMARGIN",
            new_y="NEXT",
            align="C"
        )

        pdf.ln(5)

        pdf.set_font("Hindi", size=11)

        pdf.cell(
            0,
            8,
            text="Academic Text Translator using NLP",
            new_x="LMARGIN",
            new_y="NEXT",
            align="C"
        )

        pdf.ln(8)

    # --------------------------------------------------------
    # CONTENT
    # --------------------------------------------------------

    for section in sections:

        kind = section["type"]

        text = section["text"]

        size, bold, color, before, after = pdf_style(section)

        pdf.ln(before)

        pdf.set_text_color(*color)

        pdf.set_font(
            "Hindi",
            style="B" if (bold and has_bold) else "",
            size=size
        )

        line_height = size * 0.62

        # ---------------- TITLE / SUBTITLE (centered)

        if kind in ("title", "subtitle"):

            pdf.multi_cell(
                0,
                line_height + 2,
                text=text,
                align="C",
                new_x="LMARGIN",
                new_y="NEXT"
            )

        # ---------------- HEADING

        elif kind == "heading":

            # keep a heading together with the text after it
            if pdf.get_y() > pdf.h - 45:

                pdf.add_page()

            pdf.multi_cell(
                0,
                line_height + 2,
                text=text,
                new_x="LMARGIN",
                new_y="NEXT"
            )

        # ---------------- BULLET / NUMBER (hanging indent)

        elif kind in ("bullet", "number"):

            marker = "•" if kind == "bullet" else section.get("number", "")

            pdf.set_x(pdf.l_margin + 4)

            pdf.cell(7, line_height, text=marker)

            pdf.multi_cell(
                pdf.w - pdf.r_margin - pdf.get_x(),
                line_height,
                text=text,
                new_x="LMARGIN",
                new_y="NEXT"
            )

        # ---------------- PARAGRAPH

        else:

            pdf.multi_cell(
                0,
                line_height,
                text=text,
                new_x="LMARGIN",
                new_y="NEXT"
            )

        pdf.ln(after + (2 if kind == "paragraph" else 0))

    pdf.set_text_color(0, 0, 0)

    pdf.output(output_path)

    return output_path


# ============================================================
# GENERATE DOCX
# ============================================================

def style_run(
    run,
    size=11,
    bold=False,
    color=None
):

    from docx.shared import RGBColor
    from docx.oxml.ns import qn

    run.font.size = Pt(size)

    run.font.bold = bold

    run.font.name = "Nirmala UI"

    # Make Word use Nirmala UI for Devanagari too.
    r_pr = run._element.get_or_add_rPr()

    r_fonts = r_pr.find(qn("w:rFonts"))

    if r_fonts is None:

        r_fonts = r_pr.makeelement(qn("w:rFonts"), {})

        r_pr.append(r_fonts)

    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):

        r_fonts.set(qn(attribute), "Nirmala UI")

    if color:

        run.font.color.rgb = RGBColor(*color)


def generate_docx(
    sections,
    original_filename
):

    from docx.enum.text import WD_ALIGN_PARAGRAPH

    output_name = (
        os.path.splitext(
            secure_filename(original_filename)
        )[0]
        + "_Hindi_Translation.docx"
    )

    output_path = os.path.join(
        OUTPUT_FOLDER,
        output_name
    )

    document = Document()

    has_title = any(
        section["type"] == "title"
        for section in sections
    )

    if not has_title:

        title = document.add_paragraph()

        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        style_run(
            title.add_run("शैक्षणिक पाठ का हिंदी अनुवाद"),
            size=20,
            bold=True
        )

        banner = document.add_paragraph()

        banner.alignment = WD_ALIGN_PARAGRAPH.CENTER

        style_run(
            banner.add_run("Academic Text Translator using NLP"),
            size=11
        )

    for section in sections:

        kind = section["type"]

        text = section["text"]

        size, bold, color, _, _ = pdf_style(section)

        # ---------------- TITLE / SUBTITLE

        if kind in ("title", "subtitle"):

            paragraph = document.add_paragraph()

            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            style_run(
                paragraph.add_run(text),
                size=size,
                bold=bold,
                color=color
            )

        # ---------------- HEADING (real Word headings)

        elif kind == "heading":

            level = min(max(section.get("level", 2), 1), 3)

            heading = document.add_heading(level=level)

            style_run(
                heading.add_run(text),
                size=size,
                bold=True,
                color=color
            )

        # ---------------- BULLET

        elif kind == "bullet":

            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            style_run(paragraph.add_run(text), size=11)

        # ---------------- NUMBER (keeps original number)

        elif kind == "number":

            paragraph = document.add_paragraph()

            paragraph.paragraph_format.left_indent = Pt(24)

            paragraph.paragraph_format.first_line_indent = Pt(-18)

            style_run(
                paragraph.add_run(
                    section.get("number", "") + " " + text
                ),
                size=11
            )

        # ---------------- PARAGRAPH

        else:

            paragraph = document.add_paragraph()

            style_run(paragraph.add_run(text), size=11)

    document.save(output_path)

    return output_path


# ============================================================
# HOME PAGE
# ============================================================

@app.route("/")
def home():

    return render_template(
        "index.html"
    )


# ============================================================
# UPLOAD DOCUMENT
# ============================================================

@app.route(
    "/upload",
    methods=["POST"]
)
def upload_document():

    try:

        # ----------------------------------------------------
        # CHECK FILE
        # ----------------------------------------------------

        if "file" not in request.files:

            return jsonify({

                "success": False,

                "message":
                    "No document uploaded."

            })


        file = request.files["file"]


        if not file.filename:

            return jsonify({

                "success": False,

                "message":
                    "Please select a document."

            })


        filename = secure_filename(
            file.filename
        )


        extension = os.path.splitext(
            filename
        )[1].lower()


        # ----------------------------------------------------
        # PDF
        # ----------------------------------------------------

        if extension == ".pdf":

            sections, text = extract_sections_from_pdf(
                file.read()
            )


        # ----------------------------------------------------
        # TXT
        # ----------------------------------------------------

        elif extension == ".txt":

            text = file.read().decode(

                "utf-8",

                errors="ignore"

            )

            sections = analyze_text_structure(
                text
            )


        # ----------------------------------------------------
        # INVALID FORMAT
        # ----------------------------------------------------

        else:

            return jsonify({

                "success": False,

                "message":
                    "Only PDF and TXT files are supported."

            })


        # ----------------------------------------------------
        # CHECK TEXT
        # ----------------------------------------------------

        if not text:

            return jsonify({

                "success": False,

                "message":
                    "No readable text was found. "
                    "This may be a scanned PDF."

            })


        # ----------------------------------------------------
        # WORD COUNT
        # ----------------------------------------------------

        word_count = count_words(
            text
        )


        # ----------------------------------------------------
        # DOCUMENT SIZE LIMIT
        # ----------------------------------------------------

        if word_count > MAX_DOCUMENT_WORDS:

            return jsonify({

                "success": False,

                "message":

                    f"This document contains approximately "
                    f"{word_count:,} words. "
                    f"The current maximum is "
                    f"{MAX_DOCUMENT_WORDS:,} words."

            })


        # ----------------------------------------------------
        # ANALYZE STRUCTURE
        # ----------------------------------------------------



        # ----------------------------------------------------
        # RETURN DATA
        # ----------------------------------------------------

        return jsonify({

            "success": True,

            "filename": filename,

            "extension": extension,

            "text": text,

            "sections": sections,

            "word_count": word_count,

            "character_count": len(text)

        })


    except Exception as e:

        print(
            "Upload error:",
            e
        )


        return jsonify({

            "success": False,

            "message": str(e)

        })


# ============================================================
# TRANSLATE DOCUMENT
# ============================================================

@app.route(
    "/translate-document",
    methods=["POST"]
)
def translate_document_route():

    try:

        data = request.get_json()


        sections = data.get(

            "sections",

            []

        )


        filename = data.get(

            "filename",

            "academic_document.pdf"

        )


        # ----------------------------------------------------
        # CHECK CONTENT
        # ----------------------------------------------------

        if not sections:

            return jsonify({

                "success": False,

                "message":
                    "No document content found."

            })


        print(

            f"Translating "
            f"{len(sections)} sections..."

        )


        # ----------------------------------------------------
        # TRANSLATE
        # ----------------------------------------------------

        translated_sections = (

            translate_document(

                sections

            )

        )


        print(
            "Translation completed."
        )


        # ----------------------------------------------------
        # GENERATE PDF
        # ----------------------------------------------------

        pdf_path = generate_pdf(

            translated_sections,

            filename

        )


        print(
            "PDF generated:",
            pdf_path
        )


        # ----------------------------------------------------
        # GENERATE DOCX
        # ----------------------------------------------------

        docx_path = generate_docx(

            translated_sections,

            filename

        )


        print(
            "DOCX generated:",
            docx_path
        )


        # ----------------------------------------------------
        # CREATE TRANSLATED TEXT
        # ----------------------------------------------------

        translated_text = "\n\n".join(

            section["text"]

            for section in translated_sections

        )


        # ----------------------------------------------------
        # RETURN RESULT
        # ----------------------------------------------------

        return jsonify({

            "success": True,

            "sections":
                translated_sections,

            "translation":
                translated_text,

            "pdf_url":

                "/download/pdf/"
                + os.path.basename(
                    pdf_path
                ),

            "docx_url":

                "/download/docx/"
                + os.path.basename(
                    docx_path
                )

        })


    except Exception as e:

        print(

            "Translation error:",

            e

        )


        return jsonify({

            "success": False,

            "message": str(e)

        })


# ============================================================
# DOWNLOAD PDF
# ============================================================

@app.route(
    "/download/pdf/<filename>"
)
def download_pdf(filename):

    path = os.path.join(

        OUTPUT_FOLDER,

        secure_filename(
            filename
        )

    )


    if not os.path.exists(path):

        return (

            "File not found.",

            404

        )


    return send_file(

        path,

        as_attachment=True,

        download_name=os.path.basename(
            path
        )

    )


# ============================================================
# DOWNLOAD DOCX
# ============================================================

@app.route(
    "/download/docx/<filename>"
)
def download_docx(filename):

    path = os.path.join(

        OUTPUT_FOLDER,

        secure_filename(
            filename
        )

    )


    if not os.path.exists(path):

        return (

            "File not found.",

            404

        )


    return send_file(

        path,

        as_attachment=True,

        download_name=os.path.basename(
            path
        )

    )


# ============================================================
# RUN APPLICATION
# ============================================================

if __name__ == "__main__":

    app.run(
        debug=True
    )